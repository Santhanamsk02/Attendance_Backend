import csv
import io
import re
from typing import Dict, List, Tuple, Any, Optional
import openpyxl
from fastapi import UploadFile, HTTPException, status
from pydantic import ValidationError
from app.schemas.student import StudentCreate, ImportErrorItem

MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB

# Maps common abbreviations to descriptive names as secondary fallbacks.
DEPARTMENT_ALIASES: dict[str, str] = {
    # Information Technology
    "it": "Information Technology",
    "infotech": "Information Technology",
    "information technology": "Information Technology",
    # Computer Science
    "cse": "Computer Science Engineering",
    "cs": "Computer Science Engineering",
    "computer science": "Computer Science Engineering",
    "computer science & engineering": "Computer Science Engineering",
    "computer science and engineering": "Computer Science Engineering",
    # Electronics & Communication
    "ece": "Electronics and Communication Engineering",
    "ec": "Electronics and Communication Engineering",
    "electronics": "Electronics and Communication Engineering",
    "electronics & communication": "Electronics and Communication Engineering",
    "electronics and communication": "Electronics and Communication Engineering",
    # Mechanical Engineering
    "mech": "Mechanical Engineering",
    "me": "Mechanical Engineering",
    "mechanical": "Mechanical Engineering",
    "mechanical engineering": "Mechanical Engineering",
    # AI & Data Science
    "aids": "Artificial Intelligence & Data Science",
    "ai&ds": "Artificial Intelligence & Data Science",
    "ai": "Artificial Intelligence & Data Science",
    "data science": "Artificial Intelligence & Data Science",
    "artificial intelligence": "Artificial Intelligence & Data Science",
    "artificial intelligence & data science": "Artificial Intelligence & Data Science",
    "artificial intelligence and data science": "Artificial Intelligence & Data Science",
    # AI & Machine Learning
    "aiml": "Artificial Intelligence & Machine Learning",
    "ai&ml": "Artificial Intelligence & Machine Learning",
    "artificial intelligence & machine learning": "Artificial Intelligence & Machine Learning",
    # Computer Science Business System
    "csbs": "Computer Science Business System",
    # Civil Engineering
    "civil": "Civil Engineering",
    "ce": "Civil Engineering",
    "civil engineering": "Civil Engineering",
    # Electrical & Electronics
    "eee": "Electrical & Electronics",
    "electrical": "Electrical & Electronics",
    "electrical & electronics": "Electrical & Electronics",
    "electrical and electronics": "Electrical & Electronics",
}


def normalize_department(raw: str) -> str:
    """Translate a user-supplied department string (short code or alias) to
    the canonical full name. Returns the original value if no alias is found."""
    key = raw.strip().lower()
    return DEPARTMENT_ALIASES.get(key, raw.strip())


def build_department_lookup_map(departments: List[Any]) -> Dict[str, str]:
    """
    Builds a flexible, robust department lookup dictionary mapping any code,
    name, alphanumeric variation, and common alias to the department's string ID.
    """
    dept_map: Dict[str, str] = {}
    for d in departments:
        d_id = str(d.id)
        c_code = (d.code or "").strip().lower()
        c_name = (d.name or "").strip().lower()
        if c_code:
            dept_map[c_code] = d_id
            dept_map[re.sub(r'[^a-zA-Z0-9]', '', c_code)] = d_id
        if c_name:
            dept_map[c_name] = d_id
            dept_map[re.sub(r'[^a-zA-Z0-9]', '', c_name)] = d_id

        # Map common aliases based on code and name identity
        if c_code == 'csbs' or 'business' in c_name:
            for alias in ['csbs', 'computer science business system', 'computer science & business systems', 'computer science and business systems', 'computer science business systems']:
                dept_map[alias] = d_id
        elif c_code == 'cse' or ('computer science' in c_name and 'business' not in c_name):
            for alias in ['cse', 'cs', 'computer science', 'computer science engineering', 'computer science & engineering', 'computer science and engineering']:
                dept_map[alias] = d_id
        elif c_code in ('aids', 'ai&ds') or 'data science' in c_name or ('artificial' in c_name and 'data' in c_name) or 'aritificial' in c_name:
            for alias in ['aids', 'ai&ds', 'ai-ds', 'ai and ds', 'ai & ds', 'artificial intelligence & data science', 'artificial intelligence and data science', 'aritificial intelligence & data science', 'data science']:
                dept_map[alias] = d_id
        elif c_code in ('aiml', 'ai&ml') or 'machine learning' in c_name or ('artificial' in c_name and 'machine' in c_name):
            for alias in ['aiml', 'ai&ml', 'ai-ml', 'ai and ml', 'ai & ml', 'artificial intelligence & machine learning', 'artificial intelligence and machine learning']:
                dept_map[alias] = d_id
        elif c_code in ('ece', 'ec') or 'electronics' in c_name:
            for alias in ['ece', 'ec', 'electronics', 'electronics and communication', 'electronics & communication', 'electronics and communication engineering', 'electronics & communication engineering']:
                dept_map[alias] = d_id
        elif c_code in ('mech', 'me') or 'mechanical' in c_name:
            for alias in ['mech', 'me', 'mechanical', 'mechanical engineering', 'mechanical engeering']:
                dept_map[alias] = d_id
        elif c_code == 'it' or 'information technology' in c_name:
            for alias in ['it', 'infotech', 'information technology', 'information technology (it)']:
                dept_map[alias] = d_id
        elif c_code == 'eee' or 'electrical' in c_name:
            for alias in ['eee', 'electrical', 'electrical and electronics', 'electrical & electronics', 'electrical and electronics engineering']:
                dept_map[alias] = d_id
        elif c_code in ('civil', 'ce') or 'civil' in c_name:
            for alias in ['civil', 'ce', 'civil engineering']:
                dept_map[alias] = d_id

    # Re-apply direct codes and direct names to ensure absolute priority
    for d in departments:
        d_id = str(d.id)
        if d.code:
            dept_map[d.code.strip().lower()] = d_id
            dept_map[re.sub(r'[^a-zA-Z0-9]', '', d.code.strip().lower())] = d_id
        if d.name:
            dept_map[d.name.strip().lower()] = d_id
            dept_map[re.sub(r'[^a-zA-Z0-9]', '', d.name.strip().lower())] = d_id

    return dept_map


def resolve_department(raw_dept: str, departments_map: Dict[str, str]) -> Optional[str]:
    """
    Resolves a raw department string to a department ID using departments_map.
    Tries exact match, stripped alphanumeric match, and normalized alias matching.
    """
    if not raw_dept:
        return None
    key = raw_dept.strip().lower()
    if key in departments_map:
        return departments_map[key]
    cleaned = re.sub(r'[^a-zA-Z0-9]', '', key)
    if cleaned in departments_map:
        return departments_map[cleaned]
    norm = normalize_department(raw_dept).strip().lower()
    if norm in departments_map:
        return departments_map[norm]
    cleaned_norm = re.sub(r'[^a-zA-Z0-9]', '', norm)
    if cleaned_norm in departments_map:
        return departments_map[cleaned_norm]
    return None

REQUIRED_COLUMNS = [
    "roll_no",
    "name",
    "department",
    "section",
    "academic_year",
]

ALL_COLUMNS = [
    "roll_no",
    "name",
    "email",
    "phone",
    "department",
    "section",
    "academic_year",
    "status",
]


async def parse_uploaded_file(file: UploadFile) -> List[Dict[str, Any]]:
    """
    Reads an uploaded CSV or XLSX file and returns raw row dictionaries.
    Validates file extension and size.
    """
    filename = file.filename or ""
    extension = filename.split(".")[-1].lower() if "." in filename else ""

    if extension not in ("csv", "xlsx"):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Invalid file format. Please upload a .csv or .xlsx file.",
        )

    content = await file.read()
    if len(content) > MAX_FILE_SIZE_BYTES:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="File size exceeds the 10 MB limit.",
        )

    rows: List[Dict[str, Any]] = []

    if extension == "csv":
        # Handle UTF-8 and BOM encodings
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = content.decode("latin-1")
            except Exception:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Failed to decode CSV file. Please ensure it is UTF-8 encoded.",
                )

        reader = csv.DictReader(io.StringIO(text))
        if not reader.fieldnames:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="The uploaded CSV file is empty or missing headers.",
            )

        # Normalize header names (lowercase, stripped)
        field_map = {name.strip().lower(): name for name in reader.fieldnames if name}
        
        # Check required columns
        missing = [col for col in REQUIRED_COLUMNS if col not in field_map]
        if missing:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Missing required columns in CSV: {', '.join(missing)}",
            )

        for row in reader:
            normalized_row = {
                target_col: (row.get(field_map[target_col]) or "").strip()
                for target_col in ALL_COLUMNS
                if target_col in field_map
            }
            # Fill missing optional columns as empty string
            for col in ALL_COLUMNS:
                if col not in normalized_row:
                    normalized_row[col] = ""
            rows.append(normalized_row)

    elif extension == "xlsx":
        try:
            wb = openpyxl.load_workbook(filename=io.BytesIO(content), data_only=True)
            sheet = wb.active
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Failed to read Excel (.xlsx) file: {str(e)}",
            )

        iter_rows = list(sheet.iter_rows(values_only=True))
        if not iter_rows or len(iter_rows) < 1:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="The uploaded Excel sheet is empty.",
            )

        headers = [str(cell).strip().lower() if cell is not None else "" for cell in iter_rows[0]]
        header_index_map = {h: idx for idx, h in enumerate(headers) if h}

        missing = [col for col in REQUIRED_COLUMNS if col not in header_index_map]
        if missing:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Missing required columns in Excel sheet: {', '.join(missing)}",
            )

        for row_data in iter_rows[1:]:
            # Skip completely empty rows
            if not any(row_data):
                continue
            normalized_row = {}
            for col in ALL_COLUMNS:
                if col in header_index_map:
                    idx = header_index_map[col]
                    val = row_data[idx] if idx < len(row_data) else None
                    if val is None:
                        normalized_row[col] = ""
                    else:
                        normalized_row[col] = str(val).strip()
                else:
                    normalized_row[col] = ""
            rows.append(normalized_row)

    return rows


def validate_and_sanitize_rows(
    raw_rows: List[Dict[str, Any]],
    existing_roll_nos: set,
    existing_emails: set,
    departments_map: dict,
    batches_map: dict,
) -> Tuple[List[StudentCreate], List[ImportErrorItem]]:
    """
    Validates parsed rows, detects duplicates within file and existing DB records,
    and returns valid StudentCreate instances along with a list of row errors.
    """
    valid_students: List[StudentCreate] = []
    errors: List[ImportErrorItem] = []

    seen_roll_nos_in_file = set()
    seen_emails_in_file = set()

    for idx, row in enumerate(raw_rows, start=2):  # Row 1 is header
        # Check if the row is entirely empty
        if not any(row.values()):
            continue

        raw_roll = re.sub(r"\s+", " ", (row.get("roll_no") or "")).strip().upper()
        raw_name = (row.get("name") or "").strip()
        raw_email = (row.get("email") or "").strip().lower() or None
        raw_phone = (row.get("phone") or "").strip() or None
        raw_dept_str = (row.get("department") or "").strip()
        raw_sec = (row.get("section") or "").strip().upper()
        raw_year = (row.get("academic_year") or "").strip()
        raw_status = (row.get("status") or "").strip().lower() or "active"

        # Check required roll number
        if not raw_roll:
            errors.append(ImportErrorItem(row=idx, roll_no=None, error="Roll number is required"))
            continue

        if len(raw_roll) < 3:
            errors.append(
                ImportErrorItem(
                    row=idx,
                    roll_no=raw_roll,
                    error=f"Roll number must be at least 3 characters, got '{raw_roll}'",
                )
            )
            continue

        # In-file roll_no duplicate check
        if raw_roll in seen_roll_nos_in_file:
            errors.append(
                ImportErrorItem(
                    row=idx,
                    roll_no=raw_roll,
                    error=f"Duplicate roll number '{raw_roll}' found within the uploaded file",
                )
            )
            continue

        # Database roll_no collision check
        if raw_roll in existing_roll_nos:
            errors.append(
                ImportErrorItem(
                    row=idx,
                    roll_no=raw_roll,
                    error=f"Student with roll number '{raw_roll}' already exists in database",
                )
            )
            continue

        # In-file email duplicate check
        if raw_email:
            if raw_email in seen_emails_in_file:
                errors.append(
                    ImportErrorItem(
                        row=idx,
                        roll_no=raw_roll,
                        error=f"Duplicate email '{raw_email}' found within the uploaded file",
                    )
                )
                continue
            if raw_email in existing_emails:
                errors.append(
                    ImportErrorItem(
                        row=idx,
                        roll_no=raw_roll,
                        error=f"Student with email '{raw_email}' already exists in database",
                    )
                )
                continue

        # Status normalization
        if raw_status not in ("active", "inactive"):
            raw_status = "active"

        # Validate relational foreign keys
        dept_id = resolve_department(raw_dept_str, departments_map)
        if not dept_id:
            errors.append(
                ImportErrorItem(
                    row=idx,
                    roll_no=raw_roll,
                    error=f"Department code or name '{raw_dept_str}' does not exist in the database.",
                )
            )
            continue
            
        batch = batches_map.get((str(dept_id), raw_year.lower())) or batches_map.get((dept_id, raw_year.lower()))
        if not batch:
            errors.append(
                ImportErrorItem(
                    row=idx,
                    roll_no=raw_roll,
                    error=f"Academic Batch '{raw_year}' does not exist for the specified department. Pre-register it in the Admin Dashboard.",
                )
            )
            continue

        valid_sections = [chr(65 + i) for i in range(batch.sections_count)]
        if raw_sec not in valid_sections:
            errors.append(
                ImportErrorItem(
                    row=idx,
                    roll_no=raw_roll,
                    error=f"Invalid section '{raw_sec}'. Batch '{batch.batch_name}' only supports sections: {', '.join(valid_sections)}."
                )
            )
            continue

        # Validate with Pydantic
        try:
            student_data = StudentCreate(
                roll_no=raw_roll,
                name=raw_name,
                email=raw_email,
                phone=raw_phone,
                department_id=dept_id,
                batch_id=batch.id,
                semester=batch.current_semester,
                section=raw_sec,
                academic_year=batch.batch_name,
                status=raw_status,
            )
            valid_students.append(student_data)
            seen_roll_nos_in_file.add(raw_roll)
            if raw_email:
                seen_emails_in_file.add(raw_email)
        except ValidationError as ve:
            first_err = ve.errors()[0]
            field = first_err.get("loc", ["field"])[-1]
            msg = first_err.get("msg", "Invalid value")
            errors.append(
                ImportErrorItem(
                    row=idx,
                    roll_no=raw_roll,
                    error=f"{str(field).capitalize()}: {msg}",
                )
            )
        except Exception as ex:
            errors.append(
                ImportErrorItem(
                    row=idx,
                    roll_no=raw_roll,
                    error=str(ex),
                )
            )

    return valid_students, errors
