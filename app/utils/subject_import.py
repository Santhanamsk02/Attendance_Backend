import csv
import io
from typing import Dict, List, Tuple, Any
import openpyxl
from fastapi import UploadFile, HTTPException, status
from pydantic import ValidationError
from app.schemas.subject import SubjectCreate, ImportErrorItem

MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024

DEPARTMENT_ALIASES: dict[str, str] = {
    "it": "Information Technology",
    "infotech": "Information Technology",
    "information technology": "Information Technology",
    "cse": "Computer Science",
    "cs": "Computer Science",
    "computer science": "Computer Science",
    "computer science & engineering": "Computer Science",
    "computer science and engineering": "Computer Science",
    "ece": "Electronics & Communication",
    "ec": "Electronics & Communication",
    "electronics": "Electronics & Communication",
    "electronics & communication": "Electronics & Communication",
    "electronics and communication": "Electronics & Communication",
    "mech": "Mechanical Engineering",
    "me": "Mechanical Engineering",
    "mechanical": "Mechanical Engineering",
    "mechanical engineering": "Mechanical Engineering",
    "aids": "Artificial Intelligence & Data Science",
    "ai&ds": "Artificial Intelligence & Data Science",
    "ai": "Artificial Intelligence & Data Science",
    "data science": "Artificial Intelligence & Data Science",
    "artificial intelligence": "Artificial Intelligence & Data Science",
    "artificial intelligence & data science": "Artificial Intelligence & Data Science",
    "artificial intelligence and data science": "Artificial Intelligence & Data Science",
    "civil": "Civil Engineering",
    "ce": "Civil Engineering",
    "civil engineering": "Civil Engineering",
    "eee": "Electrical & Electronics",
    "electrical": "Electrical & Electronics",
    "electrical & electronics": "Electrical & Electronics",
    "electrical and electronics": "Electrical & Electronics",
}


def normalize_department(raw: str) -> str:
    key = raw.strip().lower()
    return DEPARTMENT_ALIASES.get(key, raw.strip())


VALID_SUBJECT_TYPES = ("theory", "practical", "laboratory", "elective", "project", "other")

REQUIRED_COLUMNS = [
    "subject_code",
    "subject_name",
    "department",
    "academic_year",
    "semester",
    "subject_type",
]

ALL_COLUMNS = [
    "subject_code",
    "subject_name",
    "subject_title",
    "department",
    "academic_year",
    "semester",
    "subject_type",
    "credits",
    "description",
    "status",
]


async def parse_subject_uploaded_file(file: UploadFile) -> List[Dict[str, Any]]:
    filename = file.filename or ""
    extension = filename.split(".")[-1].lower() if "." in filename else ""

    if extension not in ("csv", "xlsx"):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Invalid file format. Please upload a .csv or .xlsx file.",
        )

    content = await file.read()
    if len(content) > MAX_FILE_SIZE_BYTES:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="File size exceeds the 10 MB limit.",
        )

    rows: List[Dict[str, Any]] = []

    if extension == "csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = content.decode("latin-1")
            except Exception:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Failed to decode CSV file. Please ensure it is UTF-8 encoded.",
                )

        reader = csv.DictReader(io.StringIO(text))
        if not reader.fieldnames:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="The uploaded CSV file is empty or missing headers.",
            )

        field_map = {name.strip().lower(): name for name in reader.fieldnames if name}
        missing = [col for col in REQUIRED_COLUMNS if col not in field_map]
        if missing:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Missing required columns in CSV: {', '.join(missing)}",
            )

        for row in reader:
            normalized_row = {
                target_col: (row.get(field_map[target_col]) or "").strip()
                for target_col in ALL_COLUMNS
                if target_col in field_map
            }
            for col in ALL_COLUMNS:
                if col not in normalized_row:
                    normalized_row[col] = ""
            rows.append(normalized_row)

    elif extension == "xlsx":
        try:
            wb = openpyxl.load_workbook(filename=io.BytesIO(content), data_only=True)
            sheet = wb.active
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Failed to read Excel (.xlsx) file: {str(e)}",
            )

        iter_rows = list(sheet.iter_rows(values_only=True))
        if not iter_rows or len(iter_rows) < 1:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="The uploaded Excel sheet is empty.",
            )

        headers = [str(cell).strip().lower() if cell is not None else "" for cell in iter_rows[0]]
        header_index_map = {h: idx for idx, h in enumerate(headers) if h}

        missing = [col for col in REQUIRED_COLUMNS if col not in header_index_map]
        if missing:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Missing required columns in Excel sheet: {', '.join(missing)}",
            )

        for row_data in iter_rows[1:]:
            if not any(row_data):
                continue
            normalized_row = {}
            for col in ALL_COLUMNS:
                if col in header_index_map:
                    idx = header_index_map[col]
                    val = row_data[idx] if idx < len(row_data) else None
                    normalized_row[col] = str(val).strip() if val is not None else ""
                else:
                    normalized_row[col] = ""
            rows.append(normalized_row)

    return rows


def validate_and_sanitize_subject_rows(
    raw_rows: List[Dict[str, Any]],
    existing_composites: set,
    departments_map: Dict[str, str],
    batches_map: Dict[str, Any],
) -> Tuple[List[SubjectCreate], List[ImportErrorItem]]:
    valid_subjects: List[SubjectCreate] = []
    errors: List[ImportErrorItem] = []

    seen_in_file = set()

    for idx, row in enumerate(raw_rows, start=2):
        if not any(row.values()):
            continue

        raw_code = (row.get("subject_code") or "").strip().upper()
        raw_name = (row.get("subject_name") or "").strip()
        raw_title = (row.get("subject_title") or "").strip() or None
        raw_dept = (row.get("department") or "").strip().upper()
        raw_year = (row.get("academic_year") or "").strip()
        raw_sem_str = (row.get("semester") or "").strip()
        raw_type = (row.get("subject_type") or "").strip().lower()
        raw_credits_str = (row.get("credits") or "").strip()
        raw_desc = (row.get("description") or "").strip() or None
        raw_status = (row.get("status") or "").strip().lower() or "active"

        if not raw_code:
            errors.append(ImportErrorItem(row=idx, subject_code=None, error="Subject code is required"))
            continue

        if not raw_name:
            errors.append(ImportErrorItem(row=idx, subject_code=raw_code, error="Subject name is required"))
            continue

        if not raw_dept:
            errors.append(ImportErrorItem(row=idx, subject_code=raw_code, error="Department is required"))
            continue
            
        department_id = departments_map.get(raw_dept)
        if not department_id:
            errors.append(ImportErrorItem(row=idx, subject_code=raw_code, error=f"Department '{raw_dept}' does not exist in the database."))
            continue

        if not raw_year:
            errors.append(ImportErrorItem(row=idx, subject_code=raw_code, error="Academic year is required"))
            continue
            
        batch_obj = batches_map.get(department_id, {}).get(raw_year.upper())
        if not batch_obj:
            errors.append(ImportErrorItem(row=idx, subject_code=raw_code, error=f"Academic Batch year '{raw_year}' does not exist under this Department in the DB. Please configure it first."))
            continue

        # Parse semester
        try:
            raw_sem = int(raw_sem_str)
            if raw_sem < 1 or raw_sem > 8:
                raise ValueError()
        except (ValueError, TypeError):
            errors.append(ImportErrorItem(row=idx, subject_code=raw_code, error=f"Invalid semester: '{raw_sem_str}'. Must be 1-8"))
            continue

        # Validate type
        if raw_type not in VALID_SUBJECT_TYPES:
            errors.append(ImportErrorItem(row=idx, subject_code=raw_code, error=f"Invalid subject type: '{raw_type}'. Must be one of: {', '.join(VALID_SUBJECT_TYPES)}"))
            continue

        # Parse credits
        raw_credits = None
        if raw_credits_str:
            try:
                raw_credits = int(raw_credits_str)
                if raw_credits < 0 or raw_credits > 10:
                    raise ValueError()
            except (ValueError, TypeError):
                errors.append(ImportErrorItem(row=idx, subject_code=raw_code, error=f"Invalid credits: '{raw_credits_str}'. Must be 0-10"))
                continue

        # Duplicate within file check
        composite_key = (raw_code, department_id, raw_year, raw_sem)
        if composite_key in seen_in_file:
            errors.append(ImportErrorItem(row=idx, subject_code=raw_code, error=f"Duplicate subject code '{raw_code}' for this department & semester within the file"))
            continue

        # Duplicate in DB
        if composite_key in existing_composites:
            errors.append(ImportErrorItem(row=idx, subject_code=raw_code, error=f"Subject code '{raw_code}' already exists for this department & semester in the database"))
            continue

        if raw_status not in ("active", "inactive"):
            raw_status = "active"

        try:
            subject_data = SubjectCreate(
                subject_code=raw_code,
                subject_name=raw_name,
                subject_title=raw_title,
                department_id=department_id,
                academic_year=raw_year,
                semester=raw_sem,
                subject_type=raw_type,
                credits=raw_credits,
                description=raw_desc,
                status=raw_status,
            )
            valid_subjects.append(subject_data)
            seen_in_file.add(composite_key)
        except ValidationError as ve:
            first_err = ve.errors()[0]
            field = first_err.get("loc", ["field"])[-1]
            msg = first_err.get("msg", "Invalid value")
            errors.append(ImportErrorItem(row=idx, subject_code=raw_code, error=f"{str(field).capitalize()}: {msg}"))
        except Exception as ex:
            errors.append(ImportErrorItem(row=idx, subject_code=raw_code, error=str(ex)))

    return valid_subjects, errors
