import csv
import io
import re
from typing import Dict, List, Tuple, Any
import openpyxl
from fastapi import UploadFile, HTTPException, status
from pydantic import ValidationError
from app.schemas.teacher import TeacherCreate, ImportErrorItem

MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB

# Redundant hardcoded aliases removed.






REQUIRED_COLUMNS = [
    "employee_id",
    "name",
    "department",
]

ALL_COLUMNS = [
    "employee_id",
    "name",
    "email",
    "phone",
    "department",
    "designation",
    "status",
]


async def parse_teacher_uploaded_file(file: UploadFile) -> List[Dict[str, Any]]:
    """
    Reads an uploaded CSV or XLSX file and returns raw row dictionaries.
    Validates file extension and size.
    """
    filename = file.filename or ""
    extension = filename.split(".")[-1].lower() if "." in filename else ""

    if extension not in ("csv", "xlsx"):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Invalid file format. Please upload a .csv or .xlsx file.",
        )

    content = await file.read()
    if len(content) > MAX_FILE_SIZE_BYTES:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="File size exceeds the 10 MB limit.",
        )

    rows: List[Dict[str, Any]] = []

    if extension == "csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = content.decode("latin-1")
            except Exception:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Failed to decode CSV file. Please ensure it is UTF-8 encoded.",
                )

        reader = csv.DictReader(io.StringIO(text))
        if not reader.fieldnames:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="The uploaded CSV file is empty or missing headers.",
            )

        field_map = {name.strip().lower(): name for name in reader.fieldnames if name}
        
        missing = [col for col in REQUIRED_COLUMNS if col not in field_map]
        if missing:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Missing required columns in CSV: {', '.join(missing)}",
            )

        for row in reader:
            normalized_row = {
                target_col: (row.get(field_map[target_col]) or "").strip()
                for target_col in ALL_COLUMNS
                if target_col in field_map
            }
            for col in ALL_COLUMNS:
                if col not in normalized_row:
                    normalized_row[col] = ""
            rows.append(normalized_row)

    elif extension == "xlsx":
        try:
            wb = openpyxl.load_workbook(filename=io.BytesIO(content), data_only=True)
            sheet = wb.active
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Failed to read Excel (.xlsx) file: {str(e)}",
            )

        iter_rows = list(sheet.iter_rows(values_only=True))
        if not iter_rows or len(iter_rows) < 1:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="The uploaded Excel sheet is empty.",
            )

        headers = [str(cell).strip().lower() if cell is not None else "" for cell in iter_rows[0]]
        header_index_map = {h: idx for idx, h in enumerate(headers) if h}

        missing = [col for col in REQUIRED_COLUMNS if col not in header_index_map]
        if missing:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Missing required columns in Excel sheet: {', '.join(missing)}",
            )

        for row_data in iter_rows[1:]:
            if not any(row_data):
                continue
            normalized_row = {}
            for col in ALL_COLUMNS:
                if col in header_index_map:
                    idx = header_index_map[col]
                    val = row_data[idx] if idx < len(row_data) else None
                    if val is None:
                        normalized_row[col] = ""
                    else:
                        normalized_row[col] = str(val).strip()
                else:
                    normalized_row[col] = ""
            rows.append(normalized_row)

    return rows


def validate_and_sanitize_teacher_rows(
    raw_rows: List[Dict[str, Any]],
    existing_employee_ids: set,
    existing_emails: set,
    departments_map: Dict[str, Any],
) -> Tuple[List[TeacherCreate], List[ImportErrorItem]]:
    """
    Validates parsed rows, detects duplicates within file and existing DB records,
    and returns valid TeacherCreate instances along with a list of row errors.
    """
    valid_teachers: List[TeacherCreate] = []
    errors: List[ImportErrorItem] = []

    seen_employee_ids_in_file = set()
    seen_emails_in_file = set()

    for idx, row in enumerate(raw_rows, start=2):
        if not any(row.values()):
            continue

        raw_emp_id = (row.get("employee_id") or "").strip().upper()
        raw_name = (row.get("name") or "").strip()
        raw_email = (row.get("email") or "").strip().lower() or None
        raw_phone = (row.get("phone") or "").strip() or None
        raw_dept_name = (row.get("department") or "").strip().lower()
        dept_id = departments_map.get(raw_dept_name)
        
        raw_desig = (row.get("designation") or "").strip() or None
        raw_status = (row.get("status") or "").strip().lower() or "active"

        if not raw_emp_id:
            errors.append(ImportErrorItem(row=idx, employee_id=None, error="Employee ID is required"))
            continue

        if len(raw_emp_id) < 2:
            errors.append(
                ImportErrorItem(
                    row=idx,
                    employee_id=raw_emp_id,
                    error=f"Employee ID must be at least 2 characters, got '{raw_emp_id}'",
                )
            )
            continue

        if raw_emp_id in seen_employee_ids_in_file:
            errors.append(
                ImportErrorItem(
                    row=idx,
                    employee_id=raw_emp_id,
                    error=f"Duplicate employee ID '{raw_emp_id}' found within the uploaded file",
                )
            )
            continue

        if raw_emp_id in existing_employee_ids:
            errors.append(
                ImportErrorItem(
                    row=idx,
                    employee_id=raw_emp_id,
                    error=f"Teacher with employee ID '{raw_emp_id}' already exists in database",
                )
            )
            continue

        if raw_email:
            if raw_email in seen_emails_in_file:
                errors.append(
                    ImportErrorItem(
                        row=idx,
                        employee_id=raw_emp_id,
                        error=f"Duplicate email '{raw_email}' found within the uploaded file",
                    )
                )
                continue
            if raw_email in existing_emails:
                errors.append(
                    ImportErrorItem(
                        row=idx,
                        employee_id=raw_emp_id,
                        error=f"Teacher with email '{raw_email}' already exists in database",
                    )
                )
                continue

        if raw_status not in ("active", "inactive"):
            raw_status = "active"

        if not dept_id:
            errors.append(
                ImportErrorItem(
                    row=idx,
                    employee_id=raw_emp_id,
                    error=f"Unrecognized department '{raw_dept_name}'. Ensure it matches the system configuration exactly."
                )
            )
            continue

        try:
            teacher_data = TeacherCreate(
                employee_id=raw_emp_id,
                name=raw_name,
                email=raw_email,
                phone=raw_phone,
                department_id=dept_id,
                designation=raw_desig,
                status=raw_status,
            )
            valid_teachers.append(teacher_data)
            seen_employee_ids_in_file.add(raw_emp_id)
            if raw_email:
                seen_emails_in_file.add(raw_email)
        except ValidationError as ve:
            first_err = ve.errors()[0]
            field = first_err.get("loc", ["field"])[-1]
            msg = first_err.get("msg", "Invalid value")
            errors.append(
                ImportErrorItem(
                    row=idx,
                    employee_id=raw_emp_id,
                    error=f"{str(field).capitalize()}: {msg}",
                )
            )
        except Exception as ex:
            errors.append(
                ImportErrorItem(
                    row=idx,
                    employee_id=raw_emp_id,
                    error=str(ex),
                )
            )

    return valid_teachers, errors
